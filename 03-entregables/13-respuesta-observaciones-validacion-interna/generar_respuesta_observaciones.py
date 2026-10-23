from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).with_name("RESPUESTA-OBSERVACIONES-validacion-interna-CyberFlow.xlsx")

NAVY = "17324D"
BLUE = "1F4E78"
TEAL = "0F6B78"
CYAN = "DDEBF7"
LIGHT_BLUE = "EAF3F8"
PALE_TEAL = "E2F0D9"
PALE_YELLOW = "FFF2CC"
PALE_RED = "FCE4D6"
PALE_GRAY = "F2F2F2"
WHITE = "FFFFFF"
BLACK = "1F1F1F"
INPUT_BLUE = "0000FF"
GREEN = "548235"
ORANGE = "C65911"
RED = "C00000"

THIN_GRAY = Side(style="thin", color="D9E1F2")
MEDIUM_BLUE = Side(style="medium", color=BLUE)


def observation(
    obs_id: str,
    kind: str,
    validator: str,
    role: str,
    minute: str,
    topic: str,
    finding: str,
    quote: str,
    answer_given: str,
    technical_check: str,
    proposed_answer: str,
    action: str,
    location: str,
    priority: str,
    status: str = "En propuesta",
) -> dict[str, str]:
    return {
        "ID": obs_id,
        "Tipo": kind,
        "Validador": validator,
        "Rol": role,
        "Minuto": minute,
        "Tema": topic,
        "Observación / recomendación": finding,
        "Cita representativa": quote,
        "Respuesta dada en sesión": answer_given,
        "Validación técnica posterior": technical_check,
        "Respuesta propuesta al evaluador": proposed_answer,
        "Acción / cambio sugerido": action,
        "Ubicación sugerida": location,
        "Prioridad": priority,
        "Estado": status,
        "Cambio aplicado": "",
        "Evidencia": "",
        "Responsable": "",
        "Fecha compromiso": "",
        "Fecha cierre": "",
    }


OBSERVATIONS = [
    observation(
        "CF-VI-01",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:33:20",
        "Rol de pfSense",
        "La arquitectura presenta pfSense como firewall, pero en la explicación se observa principalmente enrutamiento inter-VLAN. Debe precisarse qué función de seguridad cumple y qué políticas aplica.",
        "¿Por qué un firewall y no un router simple para la comunicación?",
        "Se indicó que pfSense comunica las VLAN y que algunas políticas limitan el alcance de los ataques, sin detallar las reglas.",
        "La arquitectura vigente ubica pfSense en la ruta y separa la copia SPAN del tráfico de producción. La denominación correcta depende de documentar las políticas realmente aplicadas.",
        "Se acepta la observación. En el documento se describirá como gateway pfSense para enrutamiento inter-VLAN y, solo donde exista evidencia de reglas, como firewall. Se listarán las políticas usadas para segmentación y control del laboratorio.",
        "Corregir el diagrama y el texto; incorporar una tabla de interfaces, VLAN, ruta y políticas. Si no hay evidencia de filtrado, usar 'router/gateway' en ese contexto.",
        "Arquitectura, metodología y diagrama VI-1",
        "Alta",
    ),
    observation(
        "CF-VI-02",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:34:52",
        "Uso de StandardScaler",
        "Justificar por qué se usa StandardScaler y revisar su sensibilidad a valores extremos; no asumir que Isolation Forest requiere escalado por ser no supervisado.",
        "¿Por qué esa normalización del StandardScaler?",
        "Se respondió que las variables tienen escalas muy distintas y que el escalado ayuda a uniformizarlas y fijar el umbral.",
        "La implementación desplegada empaqueta StandardScaler e Isolation Forest en un Pipeline y ajusta el scaler solo con train. Isolation Forest no es un método de distancia; por ello, la defensa debe ser experimental, no una necesidad teórica.",
        "La observación es válida. Se reemplazará la justificación incorrecta por una comparación con y sin escalado. El scaler se mantendrá únicamente si la ablación muestra una ventaja reproducible; de lo contrario, se retirará o se reportará como decisión del pipeline, no como requisito del algoritmo.",
        "Añadir ablación IF escalado vs. no escalado, documentar ajuste solo con train y discutir StandardScaler frente a alternativas robustas.",
        "Metodología de modelado, ficha del modelo y resultados",
        "Alta",
    ),
    observation(
        "CF-VI-03",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:38:04",
        "Alcance por capas",
        "La introducción sugiere análisis integral de red, pero la evaluación mostrada corresponde principalmente a capas superiores. El alcance debe anunciarse desde el inicio.",
        "Al inicio parecía que ibas a hacer un análisis completo de red.",
        "Se aclaró que se recolectan 31 variables, pero el modelo vigente usa 28 de capas 3, 4 y 7; las tres variables de capa 2 todavía no se puntúan.",
        "El alcance verificable es análisis de 28 variables L3/L4/L7 sobre IPv4. La capa 2 participa en la deduplicación del espejo, pero no en el scoring del modelo vigente.",
        "Se acepta. La tesis y la exposición declararán desde el problema y los objetivos que CyberFlow v2 analiza comportamiento L3/L4/L7 sobre IPv4. La cobertura L2 se presentará como extensión no desplegada.",
        "Uniformar título, problema, objetivos, arquitectura, metodología, resultados y limitaciones con la misma declaración de alcance.",
        "Resumen, introducción, objetivos, metodología y presentación",
        "Crítica",
    ),
    observation(
        "CF-VI-04",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:40:09",
        "Cobertura IPv6",
        "El modelo solo procesa IPv4. Debe declararse la limitación y considerar IPv6 como ampliación para evitar una zona no observada.",
        "El modelo solo trabaja para IPv4, ¿no? Vean el tema de escalamiento a IPv6.",
        "Se confirmó que el alcance actual es únicamente IPv4.",
        "El extractor y el esquema vigentes están definidos para IPv4; IPv6 no debe presentarse como cubierto ni inferirse a partir de las variables de capas superiores.",
        "Se acepta. Se declarará IPv4 como límite explícito y se propondrá una extensión IPv6 con normalización de direcciones, extracción equivalente de features, nueva línea base y recalibración antes de afirmar cobertura.",
        "Agregar limitación, amenaza a la validez y plan de trabajo futuro para IPv6; no reutilizar el umbral vigente sin validación.",
        "Alcance, limitaciones y trabajos futuros",
        "Alta",
    ),
    observation(
        "CF-VI-05",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:42:25",
        "Comparación de modelos",
        "Mostrar que Isolation Forest fue seleccionado después de comparar alternativas y explicar el criterio de selección.",
        "¿Solo has probado ese modelo o has probado algún otro modelo con el cual te puedas comparar?",
        "Se indicó que sí existe una tabla comparativa, pero no se mostró con claridad durante la exposición.",
        "Existen candidatos congelados y scripts reproducibles de comparación. La selección debe mostrar Isolation Forest, OCSVM y LOF bajo el mismo protocolo, con ROC-AUC, PR-AUC, FPR, estabilidad y limitaciones.",
        "Se acepta. Se incorporará la tabla comparativa y se explicará que Isolation Forest fue elegido por el balance medido entre detección, falsos positivos y estabilidad, no por evaluar un único algoritmo.",
        "Insertar tabla de comparación y criterios; enlazar el script compare_frozen_models_metrics.py y los artefactos congelados.",
        "Resultados de modelado y diapositiva de selección",
        "Alta",
    ),
    observation(
        "CF-VI-06",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:43:52",
        "Justificación de alpha = 0,05",
        "Explicar por qué el umbral usa alpha = 0,05 y sustentar que no es un valor arbitrario elegido después de observar la prueba.",
        "¿En qué se han basado para escoger el alfa de 0,05?",
        "Se respondió que fue recomendado para equilibrar sensibilidad y falsas alertas, sin mostrar una referencia ni el protocolo de fijación.",
        "Alpha = 0,05 fija el quinto percentil de scores normales en validación, antes de evaluar test, como presupuesto operativo objetivo cercano al 5 % de cola normal. El test retenido obtuvo 4,45 % de FPR. No es una constante universal y requiere análisis de sensibilidad.",
        "Se acepta. Se explicará alpha como una decisión de riesgo predefinida sobre validación, se mostrará la sensibilidad a otros cuantiles y se citarán el protocolo reproducible y las fuentes del algoritmo. No se afirmará que 0,05 sea universalmente óptimo.",
        "Añadir ecuación del cuantil, separación train/validation/test, análisis de sensibilidad y referencias metodológicas.",
        "Metodología de calibración, resultados y ficha del modelo",
        "Crítica",
    ),
    observation(
        "CF-VI-07",
        "Observación",
        "Ferdinand",
        "Evaluador técnico: arquitectura y modelado",
        "00:45:54",
        "Deduplicación",
        "Revisar el uso del término 'filas duplicadas'. Datos similares del mismo flujo pueden ser observaciones legítimas y no deberían eliminarse sin un criterio reproducible.",
        "Yo no creería que sea duplicado; podría ser data similar e importante.",
        "La respuesta mezcló duplicados de paquetes causados por SPAN con ventanas o filas de comportamiento similar.",
        "Hay dos problemas distintos. La deduplicación de paquetes elimina artefactos del espejo dentro de 5 ms, con campos IP/TCP y diferencia de TTL menor o igual a 1, conservando la copia de mayor TTL. Los vectores idénticos del dataset se auditan con gates; las ventanas meramente similares no deben eliminarse por parecer repetidas.",
        "Se acepta la precisión. El documento separará 'deduplicación de paquetes SPAN' de 'auditoría de vectores duplicados'. Se revisará cualquier eliminación de filas similares y se conservarán salvo que exista un criterio declarado, reproducible y sin fuga entre particiones.",
        "Corregir terminología, describir la clave y ventana de deduplicación, y documentar los gates de duplicados del dataset.",
        "Preprocesamiento, datasheet y amenazas a la validez",
        "Crítica",
    ),
    observation(
        "CF-VI-08",
        "Observación",
        "Immer Elías Cuéllar Rodríguez",
        "Evaluador: coherencia y pertinencia académica",
        "00:50:39",
        "Trazabilidad lógica",
        "Debe existir correspondencia entre lo que se promete, lo que se muestra en la práctica y lo que se concluye, especialmente respecto de las capas analizadas.",
        "Que haya una trazabilidad lógica entre lo que se habla y lo que se está mostrando después.",
        "Se agradeció la observación, sin presentar una matriz de consistencia durante la sesión.",
        "La corrección requiere una única declaración de alcance y trazabilidad objetivo-requisito-método-evidencia-resultado-limitación.",
        "Se acepta. Se incorporará una matriz de trazabilidad y se revisará que resumen, objetivos, método, demostración, resultados y conclusiones usen el mismo alcance L3/L4/L7 sobre IPv4.",
        "Crear matriz de trazabilidad y ejecutar revisión cruzada de afirmaciones en tesis, artículo, panel y presentación.",
        "Matriz de consistencia, introducción y conclusiones",
        "Crítica",
    ),
    observation(
        "CF-VI-09",
        "Observación",
        "Immer Elías Cuéllar Rodríguez",
        "Evaluador: coherencia y pertinencia académica",
        "00:51:47",
        "IPv6 como ampliación",
        "Para aumentar la relevancia publicable, explorar tráfico y ataques IPv6 sin ampliar ahora el trabajo hacia todas las capas inferiores.",
        "Sí se podría explorar los ataques o el flujo en IPv6.",
        "No se planteó un compromiso concreto durante la sesión.",
        "IPv6 es una extensión razonable, pero implica cambios de extracción, datos normales, escenarios de ataque y recalibración; no debe describirse como un cambio menor sin medirlo.",
        "Se incorpora como línea de trabajo futuro priorizada. La versión actual mantendrá el límite IPv4 y definirá un protocolo de extensión IPv6 antes de reclamar cobertura dual-stack.",
        "Añadir hoja de ruta IPv6 con dataset, features, amenazas, recalibración y evaluación dual-stack.",
        "Limitaciones, trabajos futuros y plan de mejora",
        "Media",
    ),
    observation(
        "CF-VI-10",
        "Reconocimiento",
        "Immer Elías Cuéllar Rodríguez",
        "Evaluador: coherencia y pertinencia académica",
        "00:50:10",
        "Valor del producto",
        "Se reconoce el esfuerzo invertido tanto en la detección como en la construcción de un producto de software que vuelve los resultados más intuitivos.",
        "Se ve bastante tiempo dedicado [...] también en ese producto de software.",
        "Se agradeció el reconocimiento.",
        "El reconocimiento es consistente con la existencia del panel, el motor híbrido y los artefactos reproducibles. No sustituye la evidencia de eficacia ni requiere una corrección.",
        "Se agradece el reconocimiento. Se mantendrá el énfasis en separar el aporte de producto de las métricas experimentales, evitando usar la calidad de la interfaz como evidencia de detección.",
        "Conservar como fortaleza y asegurar que el panel enlace evidencia verificable.",
        "Fortalezas del producto y discusión",
        "Baja",
        "No aplica",
    ),
    observation(
        "CF-VI-11",
        "Observación",
        "Abner Saúl Neira Torres",
        "Evaluador: operación, respuesta e integración",
        "00:53:25",
        "PERMIT / LIMIT / BLOCK",
        "Definir con precisión cuándo una anomalía recibe LIMIT o BLOCK y qué efecto técnico produce cada acción.",
        "¿En base a qué condiciones lo considera LIMIT y no BLOCK, y a qué se refiere LIMIT?",
        "La primera respuesta confundió LIMIT con alerta y bloqueo; luego se explicó que una anomalía no necesariamente es un ataque y por eso primero se limita.",
        "La regla implementada es: normal/excluido/vacío → PERMIT; anomalía del modelo → LIMIT; fuerza bruta o escaneo → BLOCK; abuso HTTP o entropía DNS → LIMIT. LIMIT aplica rate 100/second por 300 s. BLOCK usa 300/1800/3600 s según reincidencia, nunca infinito automático. La alerta es un canal de observabilidad, no una cuarta acción.",
        "Se acepta. Se corregirá la explicación para distinguir decisión, efecto y notificación. LIMIT reduce tasa como respuesta conservadora ante anomalía; BLOCK se reserva a heurísticos de mayor confianza. La acción más severa prevalece cuando coinciden señales.",
        "Añadir tabla de verdad, parámetros de nftables, caducidad, escalada, exclusiones y ejemplos por familia.",
        "Respuesta gradual, diseño de enforcement y presentación",
        "Crítica",
    ),
    observation(
        "CF-VI-12",
        "Observación",
        "Abner Saúl Neira Torres",
        "Evaluador: operación, respuesta e integración",
        "00:58:09",
        "Alcance operativo del tráfico",
        "Precisar si la solución analiza únicamente tráfico interno de la LAN o también tráfico con destino externo.",
        "¿La solución es solo a nivel de tráfico LAN o también tráfico que va con destino a la web?",
        "Se respondió que el experimento cubre tráfico interno entre las VLAN, las máquinas simuladas y el servidor.",
        "El alcance experimental es el tráfico IPv4 observable en el segmento reflejado por SPAN y atribuible a entidades monitorizadas. El enforcement se aplica en hosts Linux administrados; no equivale a protección perimetral completa.",
        "Se acepta. Se delimitará topología, punto de observación, direcciones del flujo, activos administrados y tráfico fuera de alcance. No se afirmará cobertura de Internet o del perímetro si no cruza el segmento observado.",
        "Incorporar diagrama de frontera de confianza y tabla incluido/excluido.",
        "Alcance operativo, topología y limitaciones",
        "Alta",
    ),
    observation(
        "CF-VI-13",
        "Observación",
        "Abner Saúl Neira Torres",
        "Evaluador: operación, respuesta e integración",
        "00:59:04",
        "Tiempos de detección y respuesta",
        "Separar el tiempo para detectar y clasificar del tiempo requerido para publicar y aplicar la respuesta.",
        "¿Cuánto tiempo le toma hacer ese análisis y ya tomar una decisión?",
        "Se informó un rango de 4 a 23 s, que corresponde a detección en la batería comparativa, no a la aplicación completa del control.",
        "La ablación N=3 reporta latencia de detección mediana de 13 s [4–23]. La documentación operacional amplía el rango de detección a ~2–40 s. La publicación del feed puede tardar hasta 60 s y la respuesta total observada es ~1–2,5 min, con máximo documentado de 141 s.",
        "Se corrige la respuesta: 4–23 s describe la detección de episodios de la batería, no el bloqueo extremo a extremo. Se reportarán por separado t_detección, t_decisión/publicación y t_aplicación, indicando escenario, reloj de origen y tamaño de muestra.",
        "Crear tabla de latencias con definiciones, mediana/rango, n, familia de ataque y evidencia; alinear tesis, artículo, panel y presentación.",
        "Resultados operativos, metodología de medición y diapositiva de tiempos",
        "Crítica",
    ),
    observation(
        "CF-VI-14",
        "Observación",
        "Abner Saúl Neira Torres",
        "Evaluador: operación, respuesta e integración",
        "01:00:49",
        "Integración con firewall",
        "Evaluar integración automática con el firewall para generar reglas cuando CyberFlow identifica un equipo que origina ataques.",
        "Sería bueno hacer esa integración y que automáticamente se genere una regla en el firewall.",
        "No se presentó una integración activa con pfSense; la respuesta actual se aplica mediante agentes nftables en hosts Linux.",
        "La arquitectura vigente usa feed firmado Ed25519 y agentes pull que verifican la firma antes de aplicar nftables. La integración con pfSense no está implementada y debe figurar como mejora, con controles de seguridad y reversión.",
        "Se acepta como trabajo futuro. Se diseñará un adaptador para pfSense/API que consuma decisiones firmadas, aplique expiración, lista de exclusión, aprobación humana para escaladas y rollback; no se afirmará que ya está desplegado.",
        "Definir interfaz de integración, amenazas, permisos mínimos, expiración, auditoría y prueba en laboratorio aislado.",
        "Arquitectura futura y plan de mejora",
        "Media",
    ),
    observation(
        "CF-VI-15",
        "Autocorrección",
        "Equipo expositor",
        "Responsable de corrección interna",
        "00:54:19",
        "Terminología de respuesta",
        "La explicación oral afirmó que existían básicamente dos acciones y mezcló LIMIT, ALERT y BLOCK, lo que contradice el diseño implementado.",
        "Tenemos básicamente dos acciones [...] LIMIT y ALERT son como lo mismo.",
        "La respuesta quedó ambigua y posteriormente fue matizada con el principio de respuesta gradual.",
        "El código y la metodología validan tres decisiones: PERMIT, LIMIT y BLOCK. ALERT corresponde a notificación/observabilidad, no a una acción de enforcement.",
        "Se corregirá el guion y todo material donde ALERT aparezca como acción equivalente. La demostración mostrará el efecto exacto de LIMIT y BLOCK en nftables.",
        "Actualizar guion, presentación, glosario, panel y preguntas frecuentes.",
        "Material de exposición y documentación de operación",
        "Alta",
    ),
    observation(
        "CF-VI-16",
        "Autocorrección",
        "Equipo expositor",
        "Responsable de corrección interna",
        "01:00:17",
        "Interpretación de latencia",
        "La respuesta oral pudo interpretarse como si 4–23 s fuera el tiempo total de respuesta, aunque solo corresponde a la detección de la batería comparativa.",
        "Entre cuatro segundos y 23 segundos, que es el tiempo de detección de los ataques.",
        "Se mencionó correctamente 'detección', pero no se añadió el tiempo de publicación y aplicación solicitado por el contexto operativo.",
        "Los valores son compatibles si se etiquetan: 13 s [4–23] para detección en la ablación; ~2–40 s como rango operativo documentado; hasta 60 s por publicación; ~1–2,5 min extremo a extremo, hasta 141 s.",
        "Se reemplazará cualquier frase genérica de 'tiempo de respuesta' por tres relojes definidos y se conservará la unidad experimental de cada cifra.",
        "Corregir textos, figuras y tarjetas que mezclen detección con enforcement.",
        "Resultados, presentación y panel",
        "Crítica",
    ),
    observation(
        "CF-VI-17",
        "Compromiso",
        "Nemias Saboya",
        "Moderador y seguimiento académico",
        "01:01:55",
        "Atención punto por punto",
        "Consolidar las recomendaciones, responderlas individualmente y reflejarlas en una nueva versión del trabajo.",
        "Vamos a tomar en cuenta punto por punto, vamos a hacer un resumen y considerarlo para una nueva versión.",
        "El equipo aceptó elaborar el resumen y aplicar las mejoras.",
        "Este libro materializa el registro de observaciones y permite cerrar cada punto con cambio, evidencia, responsable y fecha.",
        "Se entregará una matriz de levantamiento con estado verificable; ninguna observación se marcará como aplicada sin evidencia enlazada.",
        "Completar las columnas editables y revisar el cierre con asesor y evaluadores.",
        "Acta de validación y control de cambios",
        "Alta",
    ),
    observation(
        "CF-VI-18",
        "Compromiso",
        "Nemias Saboya",
        "Moderador y seguimiento académico",
        "00:02:27",
        "Socialización con asesor técnico",
        "La grabación y las recomendaciones deben compartirse con el profesor Fernando, asesor técnico ausente por motivos de salud, para que participe en la réplica de las mejoras.",
        "Como está quedando grabado, después escuchará para que pueda replicar las recomendaciones.",
        "Se dejó constancia de que la ausencia correspondía al asesor Fernando, distinto del evaluador Ferdinand.",
        "Se registrará la revisión del asesor técnico como parte del cierre, evitando confundir su identidad con la del evaluador Ferdinand.",
        "Se atenderá el compromiso: el profesor Fernando recibirá la grabación, esta matriz y las evidencias para revisar las correcciones técnicas antes del cierre.",
        "Compartir grabación, matriz y evidencias; registrar fecha de revisión y comentarios del asesor.",
        "Acta, participantes y control de revisión",
        "Media",
    ),
]


def set_default_font(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                font = copy(cell.font)
                font.name = "Arial"
                font.scheme = None
                cell.font = font


def style_title(sheet, title: str, subtitle: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=end_column)
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(name="Arial", size=20, bold=True, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 16
    for col in range(1, end_column + 1):
        sheet.cell(1, col).fill = PatternFill("solid", fgColor=NAVY)
        sheet.cell(2, col).fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=end_column)
    subtitle_cell = sheet.cell(3, 1, subtitle)
    subtitle_cell.font = Font(name="Arial", size=10, italic=True, color=BLUE)
    subtitle_cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 28


def style_section(cell, text: str) -> None:
    cell.value = text
    cell.font = Font(name="Arial", size=12, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=TEAL)
    cell.alignment = Alignment(vertical="center")


def build_readme(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "LEEME"
    style_title(
        sheet,
        "Matriz de respuesta a observaciones",
        "Validación interna técnica de CyberFlow | Fuente principal: grabación y transcripción de Teams | Elaborado: 15/10/2026",
        8,
    )
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A6"

    sheet.merge_cells("A5:H5")
    style_section(sheet["A5"], "Propósito y uso")
    instructions = [
        ("1", "Filtrar la hoja OBSERVACIONES por validador, prioridad o estado."),
        ("2", "Revisar la respuesta propuesta y ajustar su redacción antes de enviarla."),
        ("3", "Registrar únicamente trabajo real en las columnas amarillas: Estado, Cambio aplicado, Evidencia, Responsable y fechas."),
        ("4", "No marcar una observación como Aplicado sin indicar evidencia verificable."),
        ("5", "La hoja RESUMEN se actualiza mediante fórmulas cuando cambian los estados."),
    ]
    for index, (number, text) in enumerate(instructions, start=6):
        sheet.cell(index, 1, number)
        sheet.cell(index, 1).font = Font(name="Arial", size=11, bold=True, color=WHITE)
        sheet.cell(index, 1).fill = PatternFill("solid", fgColor=BLUE)
        sheet.cell(index, 1).alignment = Alignment(horizontal="center", vertical="center")
        sheet.merge_cells(start_row=index, start_column=2, end_row=index, end_column=8)
        sheet.cell(index, 2, text)
        sheet.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.cell(index, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE if index % 2 == 0 else WHITE)
        sheet.row_dimensions[index].height = 28

    sheet.merge_cells("A12:H12")
    style_section(sheet["A12"], "Leyenda de edición")
    legend = [
        ("Celda amarilla / texto azul", "Campo que debe completar o actualizar el equipo."),
        ("En propuesta", "La respuesta está redactada, pero el cambio todavía no ha sido ejecutado."),
        ("Aplicado", "El cambio se ejecutó y tiene evidencia enlazada."),
        ("Descartado (justificado)", "No se aplicará; la justificación debe quedar en Cambio aplicado y Evidencia."),
        ("No aplica", "Reconocimiento o punto que no requiere modificación."),
    ]
    for row, (label, meaning) in enumerate(legend, start=13):
        sheet.cell(row, 1, label)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=INPUT_BLUE if row == 13 else BLACK)
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_YELLOW if row == 13 else PALE_GRAY)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        sheet.cell(row, 3, meaning)
        sheet.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 27

    sheet.merge_cells("A19:H19")
    style_section(sheet["A19"], "Ejemplo de llenado")
    example_headers = ["Estado", "Cambio aplicado", "Evidencia", "Responsable", "Fecha compromiso", "Fecha cierre"]
    example_values = [
        "Aplicado",
        "Se añadió la tabla de políticas de pfSense en la sección 3.2.",
        "tesis-v4.pdf, p. 41; commit abc123",
        "Rubén Mark",
        date(2026, 10, 20),
        date(2026, 10, 19),
    ]
    for col, header in enumerate(example_headers, start=1):
        sheet.cell(20, col, header)
        sheet.cell(20, col).font = Font(name="Arial", size=9, bold=True, color=WHITE)
        sheet.cell(20, col).fill = PatternFill("solid", fgColor=BLUE)
        sheet.cell(20, col).alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        sheet.cell(21, col, example_values[col - 1])
        sheet.cell(21, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
        sheet.cell(21, col).font = Font(name="Arial", size=9, color=INPUT_BLUE)
        sheet.cell(21, col).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("G20:H20")
    sheet["G20"] = "Nota"
    sheet["G20"].font = Font(name="Arial", size=9, bold=True, color=WHITE)
    sheet["G20"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["G20"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("G21:H21")
    sheet["G21"] = "El ejemplo está solo en esta hoja; no forma parte del conteo de observaciones."
    sheet["G21"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    sheet["G21"].font = Font(name="Arial", size=9, color=INPUT_BLUE)
    sheet["G21"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[20].height = 32
    sheet.row_dimensions[21].height = 58
    for col in range(1, 7):
        sheet.cell(21, col).number_format = "dd/mm/yyyy" if col >= 5 else "General"

    sheet.merge_cells("A23:H23")
    style_section(sheet["A23"], "Supuestos y límites de esta matriz")
    notes = [
        "Los minutos corresponden al tiempo absoluto de la grabación de Teams.",
        "Las citas fueron limpiadas solo para legibilidad; la grabación y el VTT conservan la fuente primaria.",
        "La latencia 4–23 s pertenece a detección en la batería comparativa; no representa el enforcement extremo a extremo.",
        "Las respuestas propuestas se basan en código y documentación revisados, pero deben ser aprobadas por el equipo y el asesor antes de enviarse.",
    ]
    for row, note in enumerate(notes, start=24):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        sheet.cell(row, 1, f"• {note}")
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)
        sheet.row_dimensions[row].height = 30

    widths = {"A": 21, "B": 24, "C": 25, "D": 25, "E": 18, "F": 18, "G": 18, "H": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = "A20:H21"
    sheet.print_area = "A1:H27"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"


def build_observations(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("OBSERVACIONES")
    style_title(
        sheet,
        "Observaciones, respuestas y control de cambios",
        "Las columnas amarillas son editables. 'Cambio aplicado', evidencia, responsable y fechas se entregan en blanco para el seguimiento del equipo.",
        20,
    )
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "G8"
    sheet.auto_filter.ref = f"A7:T{7 + len(OBSERVATIONS)}"

    sheet.merge_cells("A4:T4")
    sheet["A4"] = "Criterio de cierre: una fila solo pasa a Aplicado cuando el cambio está descrito y la evidencia permite verificarlo."
    sheet["A4"].font = Font(name="Arial", size=10, bold=True, color=BLUE)
    sheet["A4"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    sheet["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[4].height = 30

    sheet.merge_cells("A5:T5")
    sheet["A5"] = "Fuente: Validación interna técnica de CyberFlow (grabación Teams, 15/10/2026), VTT y verificación posterior de código/documentación."
    sheet["A5"].font = Font(name="Arial", size=9, italic=True, color=BLACK)
    sheet["A5"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[5].height = 25

    headers = list(OBSERVATIONS[0].keys())
    header_row = 7
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, col, header)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=MEDIUM_BLUE, bottom=MEDIUM_BLUE, left=THIN_GRAY, right=THIN_GRAY)
    sheet.row_dimensions[header_row].height = 46

    editable_columns = {15, 16, 17, 18, 19, 20}
    for row_offset, item in enumerate(OBSERVATIONS, start=1):
        row = header_row + row_offset
        for col, header in enumerate(headers, start=1):
            value = item[header]
            cell = sheet.cell(row, col, value)
            cell.font = Font(name="Arial", size=9, color=INPUT_BLUE if col in editable_columns else BLACK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
            if col in editable_columns:
                cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
            elif row_offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FBFD")
        sheet.cell(row, 1).font = Font(name="Arial", size=9, bold=True, color=BLUE)
        sheet.cell(row, 5).alignment = Alignment(horizontal="center", vertical="top")
        sheet.cell(row, 14).alignment = Alignment(horizontal="center", vertical="top")
        sheet.cell(row, 15).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        sheet.cell(row, 19).number_format = "dd/mm/yyyy"
        sheet.cell(row, 20).number_format = "dd/mm/yyyy"
        sheet.row_dimensions[row].height = 145

    last_row = header_row + len(OBSERVATIONS)
    state_validation = DataValidation(type="list", formula1="=ListaEstados", allow_blank=False)
    priority_validation = DataValidation(type="list", formula1="=ListaPrioridades", allow_blank=False)
    sheet.add_data_validation(state_validation)
    sheet.add_data_validation(priority_validation)
    state_validation.add(f"O8:O200")
    priority_validation.add(f"N8:N200")

    for status, color in [
        ("Aplicado", "C6E0B4"),
        ("Por aplicar", "FCE4D6"),
        ("En propuesta", "FFF2CC"),
        ("Descartado (justificado)", "D9E1F2"),
        ("No aplica", "E7E6E6"),
    ]:
        sheet.conditional_formatting.add(
            f"O8:O{last_row}",
            FormulaRule(formula=[f'$O8="{status}"'], fill=PatternFill("solid", fgColor=color)),
        )
    for priority, color in [("Crítica", "F4CCCC"), ("Alta", "FCE5CD"), ("Media", "FFF2CC"), ("Baja", "D9EAD3")]:
        sheet.conditional_formatting.add(
            f"N8:N{last_row}",
            FormulaRule(formula=[f'$N8="{priority}"'], fill=PatternFill("solid", fgColor=color)),
        )

    widths = [
        12, 16, 26, 27, 12, 24, 55, 45, 52, 62,
        66, 55, 38, 12, 22, 50, 40, 24, 17, 15,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    sheet.print_title_rows = "1:7"
    sheet.print_area = f"A1:T{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def build_summary(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("RESUMEN")
    style_title(
        sheet,
        "Resumen de seguimiento",
        "Indicadores calculados desde la hoja OBSERVACIONES. Se actualizan al abrir el libro en Excel o LibreOffice.",
        10,
    )
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A5:B5")
    style_section(sheet["A5"], "Indicadores")
    metrics = [
        ("Total de registros", "=COUNTA(OBSERVACIONES!$A$8:$A$200)", "Número total de observaciones, reconocimientos y compromisos."),
        ("Prioridad crítica o alta", '=COUNTIF(OBSERVACIONES!$N$8:$N$200,"Crítica")+COUNTIF(OBSERVACIONES!$N$8:$N$200,"Alta")', "Puntos que requieren atención prioritaria."),
        ("Aplicados", '=COUNTIF(OBSERVACIONES!$O$8:$O$200,"Aplicado")', "Cambios ejecutados con evidencia."),
        ("Avance de cierre", '=IF(B6=0,0,(COUNTIF(OBSERVACIONES!$O$8:$O$200,"Aplicado")+COUNTIF(OBSERVACIONES!$O$8:$O$200,"Descartado (justificado)")+COUNTIF(OBSERVACIONES!$O$8:$O$200,"No aplica"))/B6)', "Aplicado + descartado justificado + no aplica, dividido entre el total."),
    ]
    for row, (label, formula, note) in enumerate(metrics, start=6):
        sheet.cell(row, 1, label)
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=BLACK)
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        sheet.cell(row, 2, formula)
        sheet.cell(row, 2).font = Font(name="Arial", size=15, bold=True, color=BLUE)
        sheet.cell(row, 2).fill = PatternFill("solid", fgColor=WHITE)
        sheet.cell(row, 2).alignment = Alignment(horizontal="center")
        sheet.cell(row, 2).number_format = "0%" if row == 9 else "0"
        sheet.cell(row, 3, note)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        sheet.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 34

    sheet.merge_cells("A12:B12")
    style_section(sheet["A12"], "Por estado")
    sheet["A13"] = "Estado"
    sheet["B13"] = "Cantidad"
    statuses = ["En propuesta", "Por aplicar", "Aplicado", "Descartado (justificado)", "No aplica"]
    for col in range(1, 3):
        sheet.cell(13, col).font = Font(name="Arial", size=9, bold=True, color=WHITE)
        sheet.cell(13, col).fill = PatternFill("solid", fgColor=BLUE)
        sheet.cell(13, col).alignment = Alignment(horizontal="center")
    for row, status in enumerate(statuses, start=14):
        sheet.cell(row, 1, status)
        sheet.cell(row, 2, f'=COUNTIF(OBSERVACIONES!$O$8:$O$200,A{row})')
        sheet.cell(row, 2).alignment = Alignment(horizontal="center")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)
        sheet.cell(row, 2).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)

    sheet.merge_cells("D12:E12")
    style_section(sheet["D12"], "Por prioridad")
    sheet["D13"] = "Prioridad"
    sheet["E13"] = "Cantidad"
    priorities = ["Crítica", "Alta", "Media", "Baja"]
    for col in range(4, 6):
        sheet.cell(13, col).font = Font(name="Arial", size=9, bold=True, color=WHITE)
        sheet.cell(13, col).fill = PatternFill("solid", fgColor=BLUE)
        sheet.cell(13, col).alignment = Alignment(horizontal="center")
    for row, priority in enumerate(priorities, start=14):
        sheet.cell(row, 4, priority)
        sheet.cell(row, 5, f'=COUNTIF(OBSERVACIONES!$N$8:$N$200,D{row})')
        sheet.cell(row, 5).alignment = Alignment(horizontal="center")
        sheet.cell(row, 4).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)
        sheet.cell(row, 5).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)

    sheet.merge_cells("G12:H12")
    style_section(sheet["G12"], "Por participante")
    sheet["G13"] = "Participante"
    sheet["H13"] = "Cantidad"
    participants = [
        "Ferdinand",
        "Immer Elías Cuéllar Rodríguez",
        "Abner Saúl Neira Torres",
        "Equipo expositor",
        "Nemias Saboya",
    ]
    for col in range(7, 9):
        sheet.cell(13, col).font = Font(name="Arial", size=9, bold=True, color=WHITE)
        sheet.cell(13, col).fill = PatternFill("solid", fgColor=BLUE)
        sheet.cell(13, col).alignment = Alignment(horizontal="center")
    for row, participant in enumerate(participants, start=14):
        sheet.cell(row, 7, participant)
        sheet.cell(row, 8, f'=COUNTIF(OBSERVACIONES!$C$8:$C$200,G{row})')
        sheet.cell(row, 8).alignment = Alignment(horizontal="center")
        sheet.cell(row, 7).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)
        sheet.cell(row, 8).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 == 0 else WHITE)

    state_chart = BarChart()
    state_chart.type = "bar"
    state_chart.style = 10
    state_chart.title = "Registros por estado"
    state_chart.y_axis.title = "Estado"
    state_chart.x_axis.title = "Cantidad"
    state_chart.height = 7
    state_chart.width = 12
    state_chart.add_data(Reference(sheet, min_col=2, min_row=13, max_row=18), titles_from_data=True)
    state_chart.set_categories(Reference(sheet, min_col=1, min_row=14, max_row=18))
    state_chart.legend = None
    sheet.add_chart(state_chart, "A21")

    participant_chart = BarChart()
    participant_chart.type = "col"
    participant_chart.style = 11
    participant_chart.title = "Registros por participante"
    participant_chart.y_axis.title = "Cantidad"
    participant_chart.height = 7
    participant_chart.width = 14
    participant_chart.add_data(Reference(sheet, min_col=8, min_row=13, max_row=18), titles_from_data=True)
    participant_chart.set_categories(Reference(sheet, min_col=7, min_row=14, max_row=18))
    participant_chart.legend = None
    sheet.add_chart(participant_chart, "F21")

    widths = {"A": 28, "B": 16, "C": 23, "D": 18, "E": 15, "F": 4, "G": 35, "H": 15, "I": 4, "J": 4}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    sheet.print_area = "A1:J38"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.orientation = "landscape"


def build_participants(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("PARTICIPANTES")
    style_title(
        sheet,
        "Participantes y fuentes",
        "Identidades verificadas con las etiquetas de hablante de Teams; la diarización local se usó únicamente como apoyo.",
        7,
    )
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"

    headers = ["Nombre", "Participación", "Enfoque principal", "Intervención", "Observaciones", "Fuente de identidad", "Nota"]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(6, col, header)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rows = [
        ["Ferdinand", "Evaluador técnico", "Arquitectura, capas, preprocesamiento y modelos", "00:33:08–00:49:38", 7, "Etiqueta Teams + diarización", "No confundir con el profesor Fernando, asesor técnico."],
        ["Immer Elías Cuéllar Rodríguez", "Evaluador", "Coherencia, pertinencia e IPv6", "00:50:03–00:52:56", 3, "Etiqueta Teams + diarización", "Incluye un reconocimiento al producto."],
        ["Abner Saúl Neira Torres", "Evaluador", "Respuesta gradual, alcance, tiempos e integración", "00:53:20–01:01:35", 4, "Etiqueta Teams + diarización", "La transcripción automática muestra también 'Amner'; se conserva el nombre de Teams."],
        ["Nemias Saboya", "Moderador", "Conducción y seguimiento punto por punto", "00:00:00–01:02:13", 2, "Etiqueta Teams + diarización", "Solicita consolidar y atender las recomendaciones."],
        ["Rubén Mark Salazar Tocas", "Equipo expositor", "Presentación, alcance, latencias y respuesta gradual", "Durante toda la sesión", 1, "Etiqueta Teams + diarización", "Responsable de parte de las respuestas orales."],
        ["Uziel Saune", "Equipo expositor", "Presentación, modelo, features y umbral", "Durante toda la sesión", 1, "Etiqueta Teams + diarización", "En Teams figura como uziel.saune."],
        ["Profesor Fernando", "Asesor técnico ausente", "Revisión técnica posterior", "Mencionado en 00:02:27", 0, "Mención de Nemias", "Ausente por salud; la grabación permite socializar las recomendaciones."],
    ]
    for row_index, values in enumerate(rows, start=7):
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor="F8FBFD" if row_index % 2 == 0 else WHITE)
            cell.border = Border(bottom=THIN_GRAY)
        sheet.row_dimensions[row_index].height = 58
    table = Table(displayName="TablaParticipantes", ref=f"A6:G{6 + len(rows)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)

    sheet.merge_cells("A16:G16")
    style_section(sheet["A16"], "Fuentes primarias y técnicas")
    sources = [
        ("Grabación", "Validación interna grabada en Microsoft Teams, duración aproximada 01:02:18."),
        ("Transcripción nominal", "Validación interna técnica de CyberFlow.vtt y versión DOCX exportadas por Teams."),
        ("Transcripción local", "Whisper small sobre la ronda de evaluadores; usada para contraste, no para sustituir el VTT."),
        ("Diarización", "Pyannote, seis voces; mapeo contrastado con las etiquetas nominales de Teams."),
        ("Verificación técnica", "Código desplegado, metodología de 7 fases, evidencias de enforcement y resultados comparativos."),
    ]
    for row, (source_type, description) in enumerate(sources, start=17):
        sheet.cell(row, 1, source_type)
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=BLUE)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        sheet.cell(row, 2, description)
        sheet.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 else WHITE)
        sheet.cell(row, 2).fill = PatternFill("solid", fgColor=PALE_GRAY if row % 2 else WHITE)
        sheet.row_dimensions[row].height = 30

    widths = {"A": 31, "B": 24, "C": 42, "D": 23, "E": 15, "F": 31, "G": 43}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.print_area = "A1:G21"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"


def build_catalogs(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("CATALOGOS")
    sheet["A1"] = "Estados"
    for row, value in enumerate(["En propuesta", "Por aplicar", "Aplicado", "Descartado (justificado)", "No aplica"], start=2):
        sheet.cell(row, 1, value)
    sheet["B1"] = "Prioridades"
    for row, value in enumerate(["Crítica", "Alta", "Media", "Baja"], start=2):
        sheet.cell(row, 2, value)
    workbook.defined_names.add(DefinedName("ListaEstados", attr_text="CATALOGOS!$A$2:$A$6"))
    workbook.defined_names.add(DefinedName("ListaPrioridades", attr_text="CATALOGOS!$B$2:$B$5"))
    sheet.sheet_state = "hidden"


def finalize(workbook: Workbook) -> None:
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.title = "Respuesta a observaciones de validación interna de CyberFlow"
    workbook.properties.subject = "Matriz de levantamiento, respuesta propuesta y seguimiento"
    workbook.properties.creator = "Equipo CyberFlow"
    workbook.properties.description = "Observaciones de la validación interna técnica y control de cambios."

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.oddHeader.center.text = "CyberFlow | Validación interna"
        sheet.oddHeader.center.size = 8
        sheet.oddHeader.center.color = "666666"
        sheet.oddFooter.left.text = "Documento de trabajo"
        sheet.oddFooter.center.text = "Página &P de &N"
        sheet.oddFooter.right.text = "15/10/2026"
        sheet.oddFooter.left.size = 8
        sheet.oddFooter.center.size = 8
        sheet.oddFooter.right.size = 8
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != "Arial":
                    cell.font = Font(
                        name="Arial",
                        size=cell.font.sz or 10,
                        bold=cell.font.bold,
                        italic=cell.font.italic,
                        color=cell.font.color,
                    )


def validate_output(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["LEEME", "OBSERVACIONES", "RESUMEN", "PARTICIPANTES", "CATALOGOS"]
    assert workbook["OBSERVACIONES"].max_row == 7 + len(OBSERVATIONS)
    assert workbook["OBSERVACIONES"]["P8"].value is None
    assert workbook["OBSERVACIONES"]["Q8"].value is None
    assert workbook["OBSERVACIONES"]["R8"].value is None
    assert workbook["RESUMEN"]["B6"].data_type == "f"
    assert len(workbook["OBSERVACIONES"].tables) == 0
    assert len(workbook["PARTICIPANTES"].tables) == 1
    assert len(workbook["OBSERVACIONES"].data_validations.dataValidation) == 2


def main() -> None:
    workbook = Workbook()
    build_readme(workbook)
    build_observations(workbook)
    build_summary(workbook)
    build_participants(workbook)
    build_catalogs(workbook)
    set_default_font(workbook)
    finalize(workbook)
    workbook.save(OUTPUT)
    validate_output(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
