"""Cierre descriptivo 4.1: scores congelados, figura escalonada y bytes UTF-8/LF.

No carga modelos, no entrena, no conecta con las VMs y no ejecuta campañas.
Después de publicar: python finalizar_41.py --verify-github <commit completo>
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import math
import os
from pathlib import Path
import platform
import re
import subprocess
import urllib.parse
import urllib.request

import matplotlib
matplotlib.use("Agg")
import numpy as np
import sklearn
from sklearn.metrics import precision_recall_curve
from PIL import Image

import cerrar_pendientes_41 as core

D = core.OUT
E = Path(__file__).resolve().parent
ROOT = core.ROOT
MAN = D / "MANIFIESTO-4.1-CANONICO.json"


def text(path, value):
    Path(path).write_bytes(value.replace("\r\n", "\n").replace("\r", "\n").encode("utf-8"))


def js(path, obj):
    text(path, json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def proposals():
    from docx import Document
    from openpyxl import load_workbook
    word = E / "Articulo-IJIES-espanol_MARK_4.1-revisado.docx"
    excel = ROOT / "03-entregables/13-respuesta-observaciones-validacion-interna/RESPUESTA-OBSERVACIONES-validacion-interna-CyberFlow.xlsx"
    doc = Document(word)
    edits = [
        ("3.3.3: composición de validación", "El subconjunto de entrenamiento se emplea",
         "El subconjunto de entrenamiento se emplea para ajustar el escalador y el Isolation Forest; el de validación, para calibrar el umbral; y el de prueba, para estimar falsos positivos. La validación contiene 65 633 ventanas nominalmente normales, incluidas 80 de la entidad Kali, 17 con actividad observable. No se confirmó independientemente su condición de ataque o normalidad; la pertenencia a Kali no basta para etiquetarlas como ataques. La Tabla 3 presenta la distribución utilizada."),
        ("3.3.3: separación de fuentes", "Los datos de ataque se mantienen separados",
         "Las 84 ventanas etiquetadas como ataque de la evaluación no se utilizaron para ajustar el escalador, entrenar el IF limpio ni calcular su umbral. Esto no implica que toda la validación tenga normalidad confirmada: contiene las 80 ventanas de Kali descritas anteriormente. La evaluación se reutilizó durante el desarrollo y no constituye un test ciego nuevo."),
        ("3.3.4: StandardScaler", "Las 28 características se organizan",
         "Las 28 características se organizan en matrices numéricas con el mismo orden en entrenamiento, validación y prueba. El Pipeline evaluado conserva StandardScaler e Isolation Forest; el escalador se ajustó únicamente con las 203 967 observaciones de entrenamiento. Se describe como componente de la configuración congelada, no como requisito teórico de IF ni como una ventaja demostrada frente a una versión sin escalado. La transformación se define en la Eq. (2)."),
        ("3.3.5: calibración", "El modelo ajustado puntúa las 65 633",
         "El IF limpio puntúa las 65 633 ventanas de validación, transformadas con el escalador de entrenamiento. La configuración congelada usa α = 0,05 y el percentil 5 de los scores para obtener el umbral. La validación incluye 80 ventanas de Kali, 17 con actividad observable, sin confirmación independiente de su condición de ataque o normalidad. El procedimiento del cuantil se expresa en la Eq. (3)."),
        ("3.3.5: evitar justificación retrospectiva", "Donde θ es el umbral;",
         "Donde θ es el umbral, qα el cuantil de orden α, sᵢ el score de normalidad de la observación estandarizada i y n el número de ventanas de validación. El corte se calcula con validación; α = 0,05 se conserva como configuración operativa congelada. No se presenta como óptimo ni como respuesta a una meta de aceptación fijada de antemano que no quedó documentada. La sensibilidad posterior es descriptiva y no selecciona un nuevo corte."),
        ("3.3.5: precisión completa del corte", "El percentil 5 produce un umbral",
         "El percentil 5 de validación produce θ = −0,5693969936322356 en la escala score_samples. La regla es score_samples < θ. Con offset_ = −0,5 equivale a −0,06939699363223562 en decision_function; ambas escalas conservan la clasificación. Esta conversión corresponde al desplazamiento de salida del IF, no al escalado de características. La Eq. (4) relaciona ambas escalas."),
        ("3.3.6: ventanas sin paquetes", "La integración conserva la puntuación",
         "La integración conserva el umbral congelado y registra el origen de cada señal. En el motor en vivo, sin paquetes recientes se omite el IF, pero todavía pueden activarse reglas sobre información retrospectiva. Por tanto, una ventana sin paquetes recientes no implica automáticamente PERMIT. En la evaluación offline de 4.1 se puntúan todas las ventanas; sus resultados no equivalen a los del motor con heurísticos."),
        ("3.3.7: unidad y alcance", "La evaluación mantiene fijos el escalador",
         "La evaluación de 4.1 mantiene fijos el escalador, el IF limpio if_limpio_2026_10 y el umbral −0,5693969936322356. Comprende 65 421 ventanas normales de prueba y 84 etiquetadas como ataque, para 65 505 observaciones. La unidad primaria es la ventana de 10 s por entidad. El conjunto se reutilizó durante el desarrollo, por lo que no es un test ciego nuevo. Estos resultados corresponden al IF offline; la comparación con heurísticos y el híbrido se reporta separadamente."),
        ("3.3.7: interpretación de métricas", "Donde  son los ataques detectados;",
         "TP y FN cuentan ventanas etiquetadas como ataque detectadas y no detectadas; FP y TN cuentan ventanas normales marcadas y no marcadas. Recall = TP/(TP+FN) y precisión = TP/(TP+FP). Se reportan también F1, FPR, ROC-AUC y AP. Las métricas de ranking usan −score_samples; AP se calcula con average_precision_score, no con integración trapezoidal. La cobertura por episodio utiliza otro denominador y se presenta por separado."),
        ("3.6: familias aproximadas", "La evaluación distingue el desempeño del detector",
         "La evaluación distingue el IF offline, el detector combinado y la aplicación de respuestas. En 4.1 las etiquetas por familia son aproximadas: HTTP identifica actividad HTTP observable y DNS actividad DNS observable; la categoría residual sin HTTP/DNS se guardó históricamente como «escaneo». Ese criterio de descarte no confirma un escaneo y no se utiliza como ground truth de familia. Cada resultado conserva el modelo, umbral, unidad de análisis y denominador."),
    ]
    result = ["# Sustituciones exactas del Word — cierre descriptivo 4.1", "",
              "No se editó el archivo fuente ni se afirma haber recibido un adjunto nuevo. Se identifican los archivos disponibles que fueron revisados.", "",
              f"- Word: `{word}`", f"- SHA-256: `{core.sha256(word)}`",
              f"- Excel: `{excel}`", f"- SHA-256: `{core.sha256(excel)}`", "",
              "Las citas son texto de párrafo extraído del Word (las ecuaciones OMML pueden no aparecer en el texto plano). Al aplicar: sustituir únicamente el párrafo indicado, conservando su estilo y las ecuaciones adyacentes.", ""]
    for title, prefix, new in edits:
        matches = [p.text for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches) != 1:
            raise ValueError(f"Ancla Word no unívoca: {prefix}: {len(matches)}")
        result += [f"## {title}", "", "**Texto actual exacto:**", "", matches[0], "",
                   "**Sustituir por:**", "", new, ""]
    result += ["## Tabla 3", "",
               "Conservar los recuentos: 337 799 antes de guarda; 2 778 descartadas; entrenamiento 203 967; validación 65 633; prueba 65 421; total 335 021.", "",
               f"**Celda actual (Validación / Propósito):** {doc.tables[2].rows[4].cells[2].text}", "",
               "**Sustituir por:** Calibrar el umbral con 65 633 ventanas nominalmente normales; incluye 80 de Kali, 17 con actividad observable, sin confirmación independiente de ataque o normalidad.", "",
               "## Tabla 5 — fila «Detección del IF»", ""]
    row = doc.tables[4].rows[2]
    replacements = ["Evaluación del IF limpio offline", "65 421 ventanas normales y 84 etiquetadas como ataque; familias proxy",
                    "65 505 ventanas de 10 s por entidad; conjunto reutilizado, no test ciego nuevo",
                    "Precisión, recall, F1, FPR, ROC-AUC y AP (average_precision_score)"]
    for i, (cell, new) in enumerate(zip(row.cells, replacements), 1):
        result += [f"- Columna {i}, actual: {cell.text}", f"  - Sustituir por: {new}"]
    result += ["", "**Nota exacta para añadir bajo Tabla 5:** Los resultados de 4.1 corresponden únicamente al IF offline. Las familias son etiquetas proxy: la categoría residual sin HTTP/DNS no confirma escaneo. Las filas de comparación híbrida y respuesta tienen protocolos y denominadores propios; no heredan las métricas de 4.1.", "",
               "## Tabla 6 y pie de figura", "",
               "Presentar F1 = 0,03485; ROC-AUC = 0,92473; AP = 0,18581 (solo aumenta la precisión de presentación, sin cambiar el valor almacenado).",
               "", "**Pie propuesto:** Figura 4. (a) ROC y (b) precisión-recall escalonada del IF limpio actual y las configuraciones de comparación, sobre 65 505 ventanas. Ranking con −score_samples y precisión completa; diagonal ROC y prevalencia PR = 84/65 505. AP se calcula con average_precision_score. IF uniform e IF exact tienen scores idénticos. La evaluación es reutilizada y describe el IF offline, no el híbrido.", "",
               "## Excel: trazabilidad, sin modificar respuestas históricas", ""]
    wb = load_workbook(excel, data_only=True, read_only=True)
    ids = {"CF-VI-02", "CF-VI-05", "CF-VI-06", "CF-VI-07", "CF-VI-11"}
    for values in wb["OBSERVACIONES"].iter_rows(values_only=True):
        if values[0] in ids:
            result.append(f"- {values[0]} — {values[5]} (estado registrado: {values[14]}).")
    wb.close()
    result += ["", "La evidencia de transferencia y α no demuestra una ablación escalado/sin escalado ni valida retrospectivamente una meta de riesgo previa. No cambiar las respuestas de sesión como si esa evidencia hubiera existido entonces. F7 no se cierra con el mero cambio de nombre de la categoría residual.", ""]
    dest = E / "SUSTITUCIONES-EXACTAS-WORD-4.1.md"
    text(dest, "\n".join(result))
    return {"word": {"path": str(word), "sha256": core.sha256(word)},
            "excel": {"path": str(excel), "sha256": core.sha256(excel)}}


def finalize():
    if core.sha256(core.FULL) != core.EXPECTED_HASH:
        raise ValueError("Cambió el respaldo canónico; detener el cierre")
    rows, fields = core.load_scores()
    metrics_before = json.loads(core.METRICS.read_text(encoding="utf-8"))
    report = []
    y = np.array([int(r["label"]) for r in rows])
    for item in metrics_before["tabla7"]:
        actual = core.model_metrics(rows, f"score::{item['modelo']}", float(item["umbral"]))
        for key, value in actual.items():
            if not math.isclose(value, item[key], rel_tol=1e-12, abs_tol=1e-12):
                raise ValueError(f"ERROR NUEVO {item['modelo']}/{key}: {value} != {item[key]}")
        score = -np.array([float(r[f"score::{item['modelo']}"]) for r in rows])
        p, rec, _ = precision_recall_curve(y, score, drop_intermediate=False)
        report.append({"modelo": item["modelo"], **actual,
                       "pr_points": len(p), "all_threshold_points": True,
                       "ap_step_sum_crosscheck": float(-np.sum(np.diff(rec) * p[:-1]))})
        assert math.isclose(report[-1]["ap_step_sum_crosscheck"], actual["average_precision"], abs_tol=1e-12)
    with (D / "TABLA7.csv").open(encoding="utf-8", newline="") as f:
        table = list(csv.DictReader(f))
    for row, item in zip(table, metrics_before["tabla7"], strict=True):
        assert row["configuracion"] == item["modelo"]
        for key in ("precision", "recall", "F1", "FPR", "roc_auc", "average_precision", "umbral"):
            assert math.isclose(float(row[key]), item[key], abs_tol=1e-12, rel_tol=1e-12)
    # Únicamente redondeos de presentación y aclaración de la etiqueta residual.
    for name in ("TABLA6.md", "TABLA6.csv"):
        content = (D / name).read_text(encoding="utf-8")
        for old, new in (("0.0348", "0.03485"), ("0.9247", "0.92473"), ("0.1858", "0.18581")):
            content = re.sub(re.escape(old) + r"(?!\d)", new, content)
        content = content.replace("ventanas de ataque detectadas", "ventanas etiquetadas como ataque detectadas")
        content = content.replace("Clase positiva = ataque (84 ventanas Kali, nota M)", "Clase positiva = 84 ventanas etiquetadas como ataque (nota M)")
        content = content.replace("| escaneo |", "| Residual sin HTTP/DNS (antes «escaneo») |")
        note = "La categoría residual se asignó por ausencia de HTTP/DNS; no confirma escaneo. Las claves históricas de metricas.json y del CSV de scores se conservan como trazabilidad, sin convertirlas en ground truth."
        if name.endswith(".md") and note not in content:
            content += "\n" + note + "\n"
        text(D / name, content)
    with (D / "TABLA6.csv").open(encoding="utf-8", newline="") as f:
        table6 = {r["indicador"]: r["valor"] for r in csv.DictReader(f)}
    clean = metrics_before["tabla6_clean"]
    for key in ("TP", "FP", "TN", "FN"):
        assert int(table6[key]) == clean[key]
    for name, key, digits in (("F1", "F1", 5), ("ROC-AUC", "roc_auc", 5),
                               ("Average Precision (AP)", "average_precision", 5)):
        assert float(table6[name]) == round(clean[key], digits)
    table7_md = (D / "TABLA7.md").read_text(encoding="utf-8")
    lines = table7_md.splitlines()
    for i, line in enumerate(lines):
        if line.startswith("| IF limpio (desplegado) |"):
            for old, new in (("0.0348", "0.03485"), ("0.9247", "0.92473"), ("0.1858", "0.18581")):
                line = re.sub(re.escape(old) + r"(?!\d)", new, line)
            lines[i] = line
    text(D / "TABLA7.md", "\n".join(lines) + "\n")
    js(core.METRICS, metrics_before)
    assert json.loads(core.METRICS.read_bytes()) == metrics_before
    p3 = json.loads(core.P3.read_text(encoding="utf-8"))
    p3["nota_politica"] = (
        "Offline se puntúan todas las ventanas. En el motor vivo, packet_count_10s=0 omite el IF, "
        "pero aún pueden activarse heurísticos sobre información retrospectiva. Los 124 FP offline "
        "en ventanas sin paquetes recientes no implican 124 acciones finales PERMIT ni determinan "
        "el FPR operativo del híbrido. Se corrige esta nota interpretativa sin cambiar recuentos."
    )
    js(core.P3, p3)
    core.make_figure(rows, metrics_before)
    # Matplotlib puede emitir espacios al final en paths SVG; no alteran geometría.
    svg = D / "figura-roc-pr-v2.svg"
    text(svg, "\n".join(line.rstrip() for line in svg.read_text(encoding="utf-8").splitlines()) + "\n")
    image = Image.open(D / "figura-roc-pr-v2.png")
    assert min(image.info["dpi"]) >= 300
    input_docs = proposals()
    # Normalización de todos los textos que se incluirán en el manifiesto.
    names = ["scores-todos-candidatos-full.csv.gz", "metricas.json", "TABLA6.md", "TABLA6.csv",
             "TABLA7.md", "TABLA7.csv", "ALFA-F6.csv", "ALFA-F6.md", "p1-reconciliacion.json",
             "p3-ventanas-cero.json", "figura-roc-pr-v2.png", "figura-roc-pr-v2.svg", "figura-roc-pr-v2.pdf"]
    for name in names:
        path = D / name
        if path.suffix in (".json", ".csv", ".md"):
            text(path, path.read_text(encoding="utf-8").rstrip("\n") + "\n")
    js(D / "COMPROBACION-METRICAS-4.1.json", {
        "rows": len(rows), "normal": 65421, "positive_labeled": 84,
        "score_source_sha256": core.sha256(core.FULL), "unchanged_metrics_values": True,
        "models": report, "tables_csv_agree": True,
        "ap_method": "sklearn.metrics.average_precision_score(y, -score_samples)",
        "ap_not_trapezoidal": True, "pr_drawstyle": "steps-post", "no_point_selection": True,
        "previous_published_discrepancy": {
            "artifact_commit": "474919802e83ab8fe879fc9a3e71c309b894cd66",
            "file": "metricas.json",
            "old_manifest_sha256": "9557d73435a61bbf91cbbf2c2a5936db6547430f906aeb17ae236d244e624e62",
            "downloaded_sha256": "bf72bfb3eb08c229bf2179791d7e0e78b0af018e8141355a3487c35dfad99233",
            "downloaded_crlf_count": 0,
            "cause": "El manifiesto se generó sobre bytes locales distintos de los normalizados por Git; no cambió el contenido JSON ni las métricas."
        },
        "png": {"pixels": list(image.size), "dpi": list(image.info["dpi"])},
        "runtime": {"python": platform.python_version(), "sklearn": sklearn.__version__, "matplotlib": matplotlib.__version__},
        "reviewed_sources": input_docs,
    })
    names.append("COMPROBACION-METRICAS-4.1.json")
    for path in (E / "BASE-EDITORIAL-FINAL-4.1.md", E / "BASE-EDITORIAL-4.1.md",
                 E / "SUSTITUCIONES-EXACTAS-WORD-4.1.md",
                 E / "finalizar_41.py", E / "cerrar_pendientes_41.py", ROOT / ".gitattributes"):
        names.append(Path(os.path.relpath(path, D)).as_posix())
    assert all((D / n).exists() for n in names)
    for name in names:
        path = D / name
        if path.suffix in (".json", ".csv", ".md", ".py", ".svg") or path.name == ".gitattributes":
            text(path, path.read_text(encoding="utf-8").rstrip("\n") + "\n")
    index = E / "ENTREGABLES-FINALES-4.1.md"
    proof = D / "VERIFICACION-GITHUB-4.1.json"
    paths = [(D / n).resolve() for n in names] + [MAN, proof, index]
    lines = ["# Archivos finales — cierre 4.1", "",
             "Las rutas se listan completas. Los enlaces apuntan a main; la verificación de integridad registra el commit inmutable de los artefactos.", "",
             "| Archivo local | GitHub |", "|---|---|"]
    for path in paths:
        rel = path.relative_to(ROOT).as_posix()
        url = "https://github.com/marksato13/VF-PPI-TESIS-ORQUESTACION/blob/main/" + urllib.parse.quote(rel, safe="/")
        lines.append(f"| `{path}` | [{path.name}]({url}) |")
    lines += ["", "El informe de integridad se genera después de publicar los artefactos y se publica en un segundo commit; no se incluye en su propio manifiesto.", ""]
    text(index, "\n".join(lines))
    names.append(Path(os.path.relpath(index, D)).as_posix())
    js(MAN, {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_scores": {"file": core.FULL.name, "sha256": core.sha256(core.FULL), "rows": len(rows), "fields": fields},
        "expected_hash_match": True, "metrics_values_unchanged": True,
        "serialization": "UTF-8 sin BOM, LF; JSON indent=2 ensure_ascii=False con LF final; SHA-256 de bytes finales",
        "figure_script": "finalizar_41.py (make_figure de cerrar_pendientes_41.py)",
        "figure_method": "ROC completa y PR steps-post; todos los umbrales; AP=average_precision_score; sin integración trapezoidal",
        "alpha_source": {"file": "p1-reconciliacion.json", "n_val": 65633},
        "files": {n: core.sha256(D / n) for n in names},
        "model_or_deployment_changes": False,
        "integrity_proof": "VERIFICACION-GITHUB-4.1.json se genera después del push y no se auto-incluye en los hashes",
    })
    print(json.dumps({"checked_models": len(report), "metrics_unchanged": True, "manifest": str(MAN)}, indent=2))


def verify_github(commit):
    if len(commit) != 40 or any(c not in "0123456789abcdef" for c in commit):
        raise ValueError("Indicar commit completo e inmutable")
    base = f"https://raw.githubusercontent.com/marksato13/VF-PPI-TESIS-ORQUESTACION/{commit}/"
    def fetch(relative):
        url = base + urllib.parse.quote(relative, safe="/")
        with urllib.request.urlopen(url, timeout=90) as response:
            return response.read(), url
    rel = MAN.relative_to(ROOT).as_posix()
    manifest_bytes, manifest_url = fetch(rel)
    manifest = json.loads(manifest_bytes)
    checks = []
    for name, expected in manifest["files"].items():
        path = (D / name).resolve().relative_to(ROOT).as_posix()
        data, url = fetch(path)
        actual = hashlib.sha256(data).hexdigest()
        checks.append({"file": name, "url": url, "expected_sha256": expected,
                       "downloaded_sha256": actual, "match": expected == actual, "bytes": len(data)})
    bad = [c for c in checks if not c["match"]]
    proof = {"checked_at_utc": datetime.now(timezone.utc).isoformat(), "artifact_commit": commit,
             "manifest_url": manifest_url, "downloaded_manifest_sha256": hashlib.sha256(manifest_bytes).hexdigest(),
             "verification": "Descarga HTTP de cada archivo desde raw.githubusercontent.com con commit fijo; hashes sobre bytes recibidos, sin normalización",
             "files_checked": len(checks), "discrepancies": len(bad), "checks": checks}
    js(D / "VERIFICACION-GITHUB-4.1.json", proof)
    print(json.dumps({"files_checked": len(checks), "discrepancies": len(bad), "artifact_commit": commit}))
    if bad:
        raise SystemExit("Integridad remota fallida; no declarar cierre")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--verify-github")
    args = parser.parse_args()
    verify_github(args.verify_github) if args.verify_github else finalize()
